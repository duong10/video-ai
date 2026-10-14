#!/usr/bin/env python3

"""Tạo giọng đọc từ file kịch bản Word (.docx) và ghi mốc thời gian thật vào kịch bản.

Thư viện cần dùng đã có trong requirements.txt.

Dự án video nằm trong du-an/ và được tạo giọng qua ./tao-video du-an/<tên> (hoặc
./giao-dien), luồng đó gọi process() ở đây.

Chạy riêng công cụ này với file kịch bản nằm ở đâu cũng được:
         .venv/bin/python san_xuat/tao_tu_docx.py <thư mục chứa .docx | file .docx>
  Giọng đọc ghi vào output/ cạnh file .docx, còn file .docx được sửa lại cột thời
  gian, dòng "Thời lượng" và tiêu đề theo giọng đọc thực tế.

File Word cần có:
  - Bảng chính có cột "Cảnh / Thời gian" (dòng đầu "1. Mở đầu", dòng sau là mốc giờ,
    có thể để trống) và cột "Lời đọc". Tạo ra output/video_90s/.
  - (Không bắt buộc) bảng rút gọn có cột "Thời gian" và "Lời đọc". Tạo ra
    output/quang_cao_30s/.
  - (Không bắt buộc) dòng "Thời lượng: ... giọng vi-VN-..., tốc độ +10%" để chọn
    giọng và tốc độ chung.

Muốn chỉnh riêng từng cảnh thì đặt file tinh_chinh.toml cạnh file .docx:

    giong = "vi-VN-HoaiMyNeural"   # ghi đè giọng trong file Word
    toc_do = "+10%"                # ghi đè tốc độ chung

    [cach_doc]                     # chữ trên phụ đề = cách giọng đọc
    "8Car" = "Tám Ca"

    [video_90s.2]                  # cảnh số 2 của bảng chính
    toc_do = "+15%"
    nghi_truoc = 0.6               # khoảng nghỉ trước cảnh (giây)
    nghi_giua_cau = 0.3            # khoảng nghỉ giữa các câu (giây)

    [quang_cao_30s.1]              # dòng số 1 của bảng rút gọn
    toc_do = "+15%"
"""

import argparse
import asyncio
import re
import sys
import tomllib
import unicodedata
from pathlib import Path
from typing import Dict, List, NamedTuple, Optional, Tuple

from docx import Document
from docx.shared import Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

import giong_doc
from giong_doc import Scene, SceneTime

MAIN_NAME = "video_90s"
AD_NAME = "quang_cao_30s"


class Part(NamedTuple):
    """Một bảng kịch bản đã đọc từ file Word."""

    name: str  # tên thư mục output
    table: Table
    time_col: int
    scenes: List[Scene]
    time_cells: List[_Cell]


def slugify(text: str, max_words: int = 0) -> str:
    """ "Việc 1: Xác minh giấy tờ" -> "viec_1_xac_minh_giay_to"."""
    text = unicodedata.normalize("NFD", text.replace("đ", "d").replace("Đ", "D"))
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    words = re.findall(r"[a-z0-9]+", text)
    return "_".join(words[:max_words] if max_words else words)


def narration(cell: _Cell) -> str:
    """Lời đọc trong ô: phần nằm trong ngoặc kép, bỏ ghi chú bên ngoài."""
    text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
    first, last = text.find("“"), text.rfind("”")
    if first != -1 and last > first:
        text = text[first + 1 : last]
    return re.sub(r"\s+", " ", text.replace('"', "")).strip()


def header_index(table: Table, *names: str) -> Optional[int]:
    headers = [c.text.strip().lower() for c in table.rows[0].cells]
    for i, h in enumerate(headers):
        if any(h.startswith(n) for n in names):
            return i
    return None


def scene_options(
    tuning: dict, part: str, number: int, base_rate: str, first: bool
) -> dict:
    opts = tuning.get(part, {}).get(str(number), {})
    return {
        "rate": opts.get("toc_do", base_rate),
        "pause_before": opts.get(
            "nghi_truoc", giong_doc.LEAD_IN if first else giong_doc.SCENE_PAUSE
        ),
        "sentence_pause": opts.get("nghi_giua_cau", giong_doc.SENTENCE_PAUSE),
    }


def read_parts(doc, base_rate: str, tuning: dict) -> List[Part]:
    parts = []
    for table in doc.tables:
        text_col = header_index(table, "lời đọc")
        if text_col is None:
            continue
        scene_col = header_index(table, "cảnh")
        time_col = header_index(table, "thời gian")
        rows = [r for r in table.rows[1:] if narration(r.cells[text_col])]

        scenes, cells = [], []
        if scene_col is not None:  # bảng chính: "1. Mở đầu" + mốc giờ
            name = MAIN_NAME
            for n, row in enumerate(rows, 1):
                cell = row.cells[scene_col]
                title = cell.paragraphs[0].text.strip()
                m = re.match(r"(\d+)\.?\s*(.*)", title)
                number = int(m.group(1)) if m else n
                label = m.group(2) if m else title
                slug = f"{number:02d}_{slugify(label)}"
                opts = scene_options(tuning, name, number, base_rate, n == 1)
                scenes.append(Scene(slug, narration(row.cells[text_col]), **opts))
                cells.append(cell)
        elif time_col is not None:  # bảng rút gọn: "0 – 2.5s"
            name = AD_NAME
            for n, row in enumerate(rows, 1):
                text = narration(row.cells[text_col])
                slug = f"{n:02d}_{slugify(text, 4)}"
                opts = scene_options(tuning, name, n, base_rate, n == 1)
                scenes.append(Scene(slug, text, **opts))
                cells.append(row.cells[time_col])
        else:
            continue
        if any(p.name == name for p in parts):
            raise SystemExit(f"File Word có nhiều hơn một bảng dạng {name}")
        parts.append(Part(name, table, time_col, scenes, cells))
    return parts


# ---- Ghi mốc thời gian ngược vào file Word ----


def set_runs(paragraph: Paragraph, text: str, keep: int = 0) -> None:
    """Thay chữ của paragraph sau run thứ keep, giữ định dạng của run đầu tiên được thay."""
    runs = paragraph.runs
    if len(runs) <= keep:
        paragraph.add_run(text)
        return
    runs[keep].text = text
    for run in runs[keep + 1 :]:
        run.text = ""


def fmt_main(seconds: float) -> str:
    minutes, sec = divmod(round(seconds, 1), 60)
    return f"{int(minutes)}:{sec:04.1f}"


def fmt_ad(seconds: float) -> str:
    return "0" if round(seconds, 1) == 0 else f"{seconds:.1f}"


def fmt_duration(total: float) -> str:
    rounded = round(total)
    minutes, sec = divmod(rounded, 60)
    words = f"{minutes} phút {sec} giây" if minutes else f"{sec} giây"
    if minutes and not sec:
        words = f"{minutes} phút"
    return f"{words} ({total:.1f} giây)".replace(".", ",")


def write_times(part: Part, times: List[SceneTime]) -> List[Tuple[str, str]]:
    """Ghi mốc thật vào cột thời gian. Trả về (mốc cũ, mốc mới) của từng cảnh."""
    changes = []
    for cell, t in zip(part.time_cells, times):
        if part.name == MAIN_NAME:
            new = f"{fmt_main(t.start)} – {fmt_main(t.end)}"
            if len(cell.paragraphs) > 1:
                para = cell.paragraphs[1]
                old = para.text.strip()
                set_runs(para, new)
            else:
                old = ""
                run = cell.add_paragraph().add_run(new)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        else:
            new = f"{fmt_ad(t.start)} – {fmt_ad(t.end)}s"
            para = cell.paragraphs[0]
            old = para.text.strip()
            set_runs(para, new)
        changes.append((old, new))
    return changes


def rate_text(base_rate: str, scenes: List[Scene]) -> str:
    special = [
        f"cảnh {s.slug.split('_')[0].lstrip('0')}: {s.rate}"
        for s in scenes
        if s.rate != base_rate
    ]
    text = f"tốc độ {base_rate}"
    return f"{text} (riêng {', '.join(special)})" if special else text


def update_headings(doc, totals: Dict[str, float], voice: str, base_rate: str, scenes) -> None:
    for para in doc.paragraphs:
        text = para.text
        if text.startswith("Thời lượng:") and MAIN_NAME in totals:
            value = (
                f" {fmt_duration(totals[MAIN_NAME])}, tính theo file giọng đọc "
                f"{MAIN_NAME}.mp3 – giọng {voice}, {rate_text(base_rate, scenes)}"
            )
            # Giữ nguyên run nhãn "Thời lượng:" (in đậm), thay phần còn lại.
            label_runs = 0
            seen = ""
            for run in para.runs:
                seen += run.text
                label_runs += 1
                if "Thời lượng:" in seen:
                    break
            if seen.strip() == "Thời lượng:":
                set_runs(para, value.lstrip() if seen.endswith(" ") else value, keep=label_runs)
            else:
                set_runs(para, "Thời lượng:" + value)
            continue
        for pattern, name in (
            (r"(Kịch bản chi tiết \()\d+( giây)", MAIN_NAME),
            (r"(Bản rút gọn )\d+( giây)", AD_NAME),
        ):
            if name in totals and re.search(pattern, text):
                seconds = round(totals[name])
                new = re.sub(pattern, rf"\g<1>{seconds}\g<2>", text)
                for run in para.runs:
                    if re.search(r"\d+ giây", run.text):
                        run.text = re.sub(r"\d+( giây)", rf"{seconds}\g<1>", run.text)
                        break
                else:
                    set_runs(para, new)


# ---- Luồng chính ----


def show(path: Path) -> str:
    cwd = Path.cwd()
    return str(path.relative_to(cwd)) if path.is_relative_to(cwd) else str(path)


def find_docx(folder: Path) -> Path:
    files = [f for f in folder.glob("*.docx") if not f.name.startswith("~$")]
    if len(files) != 1:
        raise SystemExit(f"{folder} cần đúng 1 file .docx, đang có {len(files)}")
    return files[0]


async def process(
    docx_path: Path, out_dir: Optional[Path] = None, tuning: Optional[dict] = None
) -> None:
    """out_dir: nơi ghi giọng đọc (mặc định <thư mục .docx>/output);
    tuning: các giá trị tinh chỉnh (mặc định đọc tinh_chinh.toml cạnh file .docx)."""
    folder = docx_path.parent
    out_dir = out_dir or folder / "output"
    if tuning is None:
        tuning_file = folder / "tinh_chinh.toml"
        tuning = tomllib.loads(tuning_file.read_text("utf-8")) if tuning_file.exists() else {}

    doc = Document(docx_path)
    header = "\n".join(p.text for p in doc.paragraphs[:15])
    voice_match = re.search(r"\b([a-z]{2}-[A-Z]{2}-\w+Neural)\b", header)
    rate_match = re.search(r"tốc độ ([+-]\d+%)", header)
    voice = tuning.get("giong") or (voice_match.group(1) if voice_match else giong_doc.DEFAULT_VOICE)
    base_rate = tuning.get("toc_do") or (rate_match.group(1) if rate_match else giong_doc.DEFAULT_RATE)
    pronunciations = {**giong_doc.DEFAULT_PRONUNCIATIONS, **tuning.get("cach_doc", {})}

    parts = read_parts(doc, base_rate, tuning)
    if not any(p.name == MAIN_NAME for p in parts):
        raise SystemExit(
            f"{docx_path.name}: không thấy bảng có cột 'Cảnh / Thời gian' và 'Lời đọc'"
        )

    print(f"=== {folder.name if out_dir == folder / 'output' else docx_path.stem}  (giọng {voice}, tốc độ {base_rate})")
    totals = {}
    for part in parts:
        times, total = await giong_doc.build(out_dir / part.name, part.scenes, voice, pronunciations)
        totals[part.name] = total
        changes = write_times(part, times)
        print(f"\n{part.name}: {total:.1f}s")
        print(f"  {'Cảnh':<38}{'Trong file cũ':>18}{'Thực tế':>18}")
        for t, (old, new) in zip(times, changes):
            mark = "" if old == new else "  *"
            print(f"  {t.scene.slug:<38}{old or '(trống)':>18}{new:>18}{mark}")

    main_scenes = next(p.scenes for p in parts if p.name == MAIN_NAME)
    update_headings(doc, totals, voice, base_rate, main_scenes)
    doc.save(docx_path)
    print(f"\nĐã ghi mốc thời gian thật vào {show(docx_path)}")
    print(f"Giọng đọc trong {show(out_dir)}/\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("thu_muc", nargs="+", help="thư mục chứa kịch bản .docx hoặc chính file .docx")
    args = parser.parse_args()

    targets = []
    for item in args.thu_muc:
        path = Path(item).resolve()
        targets.append(path if path.suffix == ".docx" else find_docx(path))

    for target in targets:
        try:
            asyncio.run(process(target))
        except Exception:
            print(
                f"\nLỗi khi tạo giọng cho {show(target)}. Sửa xong thì chạy lại:\n"
                f"  .venv/bin/python san_xuat/tao_tu_docx.py {show(target.parent)}\n",
                file=sys.stderr,
            )
            raise


if __name__ == "__main__":
    sys.exit(main())
