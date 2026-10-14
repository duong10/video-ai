"""Đọc kịch bản Word cho bước dựng video và ghép ảnh trong thư mục dự án vào từng cảnh.

Phần lời đọc do tao_tu_docx.py xử lý. Ở đây chỉ lấy những gì video
cần: tên cảnh, mô tả hình ảnh, chữ trên màn hình và chữ to trên ảnh bìa.
"""

import re
import unicodedata
from pathlib import Path
from typing import Dict, List, NamedTuple, Optional

from docx import Document

from tao_tu_docx import header_index, narration

ROOT = Path(__file__).resolve().parent.parent

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"


class ScriptScene(NamedTuple):
    number: int
    title: str  # "Nỗi đau"
    image_note: str  # cột "Hình ảnh"
    screen_lines: List[str]  # cột "Chữ trên màn hình", mỗi đoạn một dòng
    frame_count: int  # "Cắt nhanh 4 khung" -> 4, không ghi -> 1
    frame_notes: List[str]  # các ý ① ② ③ trong cột "Hình ảnh"
    narration: str = ""  # lời đọc của cảnh


class Script(NamedTuple):
    title: str
    cover_text: Optional[str]
    scenes: List[ScriptScene]


class Images(NamedTuple):
    cover: Optional[Path]
    scenes: Dict[int, List[Path]]
    unknown: List[Path]


def nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text)


def find_docx(folder: Path) -> Path:
    files = [f for f in folder.glob("*.docx") if not f.name.startswith("~$")]
    if len(files) != 1:
        raise SystemExit(f"{folder} cần đúng 1 file kịch bản .docx, đang có {len(files)}")
    return files[0]


def last_quote(text: str) -> Optional[str]:
    quotes = re.findall(r"“([^”]+)”", text) or re.findall(r'"([^"]+)"', text)
    return quotes[-1].strip() if quotes else None


def read_script(docx_path: Path) -> Script:
    doc = Document(docx_path)
    paragraphs = [nfc(p.text.strip()) for p in doc.paragraphs if p.text.strip()]
    title = paragraphs[1] if len(paragraphs) > 1 else docx_path.stem
    cover_text = next(
        (last_quote(p) for p in paragraphs if p.lower().startswith("ảnh bìa")), None
    )

    for table in doc.tables:
        text_col = header_index(table, "lời đọc")
        scene_col = header_index(table, "cảnh")
        if text_col is None or scene_col is None:
            continue
        image_col = header_index(table, "hình ảnh")
        screen_col = header_index(table, "chữ trên màn hình", "chữ")
        scenes = []
        rows = [r for r in table.rows[1:] if narration(r.cells[text_col])]
        for n, row in enumerate(rows, 1):
            heading = nfc(row.cells[scene_col].paragraphs[0].text.strip())
            m = re.match(r"(\d+)\.?\s*(.*)", heading)
            number, label = (int(m.group(1)), m.group(2)) if m else (n, heading)

            image_lines = []
            if image_col is not None:
                image_lines = [nfc(p.text.strip()) for p in row.cells[image_col].paragraphs]
                image_lines = [line for line in image_lines if line]
            note = " ".join(image_lines)
            count = re.search(r"(\d+)\s*khung", note)
            frame_notes = [
                line[1:].strip() for line in image_lines if line[:1] in CIRCLED
            ]

            screen = []
            if screen_col is not None:
                screen = [nfc(p.text.strip()) for p in row.cells[screen_col].paragraphs]
                screen = [line for line in screen if line]

            scenes.append(
                ScriptScene(
                    number,
                    label,
                    note,
                    screen,
                    int(count.group(1)) if count else 1,
                    frame_notes,
                    nfc(narration(row.cells[text_col])),
                )
            )
        return Script(title, cover_text, scenes)
    raise SystemExit(f"{docx_path.name}: không thấy bảng có cột 'Cảnh' và 'Lời đọc'")


# ---- Ghép ảnh vào cảnh ----


def natural_key(path: Path):
    stem = nfc(path.stem).lower()
    return [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", stem)]


def scene_of(path: Path) -> Optional[int]:
    """Số cảnh của ảnh theo tên file; 0 là ảnh bìa.

    Nhận được: "Cảnh 2 · Nỗi đau 1:4 ....png", "02b_noi_lo.png", "2.png",
    "Ảnh bìa (00-anh-bia).png", "00_anh_bia.png", "anh-bia.jpg".
    """
    stem = nfc(path.stem).lower()
    if "bìa" in stem or re.search(r"(^|[^a-z])bia([^a-z]|$)", stem):
        return 0
    for pattern in (r"cảnh\s*(\d+)", r"canh[\s_-]*(\d+)", r"^(\d+)", r"\((\d+)[-_ ]"):
        m = re.search(pattern, stem)
        if m:
            return int(m.group(1))
    return None


def find_images(folder: Path) -> Images:
    cover, scenes, unknown = None, {}, []
    for path in sorted(folder.iterdir(), key=natural_key):
        if path.suffix.lower() not in IMAGE_EXT or path.name.startswith("."):
            continue
        number = scene_of(path)
        if number is None:
            unknown.append(path)
        elif number == 0:
            cover = cover or path
        else:
            scenes.setdefault(number, []).append(path)
    return Images(cover, scenes, unknown)


def is_cta(sc: ScriptScene) -> bool:
    """Cảnh kêu gọi hành động: chữ trên màn hình có hotline, số điện thoại hoặc "dùng thử"."""
    text = " ".join(sc.screen_lines).lower()
    return bool(re.search(r"hotline|\d{3,4}[. ]\d{3,4}[. ]\d{3}|dùng thử", text))


def check_images(script: Script, images: Images) -> List[str]:
    """In cảnh báo; trả về danh sách lỗi (thiếu ảnh) để dừng trước khi tạo giọng."""
    errors = []
    for sc in script.scenes:
        found = images.scenes.get(sc.number, [])
        if not found and sc is script.scenes[-1] and is_cta(sc) and len(script.scenes) > 1:
            print(f"  ! Cảnh {sc.number} (kêu gọi hành động) không có ảnh: dùng ảnh cảnh trước làm mờ làm nền")
        elif not found:
            errors.append(f"Cảnh {sc.number} · {sc.title}: thiếu ảnh. Mô tả: {sc.image_note}")
        elif sc.frame_count > 1 and len(found) != sc.frame_count:
            print(
                f"  ! Cảnh {sc.number} ghi {sc.frame_count} khung nhưng có {len(found)} ảnh"
                " - vẫn dựng với số ảnh hiện có"
            )
    numbers = {sc.number for sc in script.scenes}
    for n in sorted(set(images.scenes) - numbers):
        print(f"  ! Có ảnh cho cảnh {n} nhưng kịch bản không có cảnh này: bỏ qua")
    for path in images.unknown:
        print(f"  ! Không biết ảnh '{path.name}' thuộc cảnh nào (tên cần có số cảnh): bỏ qua")
    if not images.cover:
        print("  ! Không có ảnh bìa (tên chứa 'bìa' hoặc bắt đầu bằng 00): video mở thẳng vào cảnh 1")
    return errors
